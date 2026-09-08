"""Marketplace listing-workbook layout for parse and upload.

Offline adapters in ``ops/listing_mapping/marketplace/`` wrap this config.
Runtime upload uses the same JSON so Flipkart never inherits Amazon sheet offsets.
"""

from __future__ import annotations

import io
import json
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field

from dto.listing_config import ListingTemplateMetadata
from utils.listing_template_columns import EnumDiscovery, WorkbookLayout
from utils.listing_workbook import WorkbookKind, detect_openxml_kind

_CONFIG_PATH = Path(__file__).resolve().parent / "marketplace_listing_workbooks.json"

_MARKETPLACE_KEYS = ("AMAZON", "FLIPKART", "MYNTRA")


class MarketplaceWorkbookConfig(BaseModel):
    """Default blank-workbook layout for one marketplace."""

    model_config = ConfigDict(extra="forbid")

    sheet_name: str
    header_label_row: int = Field(ge=1)
    machine_key_row: int = Field(ge=1)
    data_start_row: int = Field(ge=1)
    enum_discovery: EnumDiscovery
    valid_values_sheet: str | None = None
    dropdown_lists_sheet: str | None = None
    data_definitions_sheet: str | None = None


@lru_cache(maxsize=1)
def _load_raw() -> dict[str, dict]:
    if not _CONFIG_PATH.is_file():
        raise FileNotFoundError(f"Missing marketplace workbook config: {_CONFIG_PATH}")
    raw = json.loads(_CONFIG_PATH.read_text(encoding="utf-8"))
    if not isinstance(raw, dict):
        raise TypeError(f"{_CONFIG_PATH} must be a JSON object keyed by marketplace id")
    return raw


def marketplace_key_from_name(name: str) -> str:
    """Map a ``marketplace.name`` row to AMAZON / FLIPKART / MYNTRA."""
    needle = name.strip().casefold()
    matches = [key for key in _MARKETPLACE_KEYS if key.casefold() in needle]
    if len(matches) == 1:
        return matches[0]
    known = ", ".join(_MARKETPLACE_KEYS)
    raise ValueError(
        f"Cannot resolve listing-workbook adapter for marketplace {name!r}. "
        f"Name must uniquely contain one of: {known}."
    )


def config_for_key(marketplace_key: str) -> MarketplaceWorkbookConfig:
    raw = _load_raw()
    entry = raw.get(marketplace_key)
    if entry is None:
        known = ", ".join(sorted(raw.keys())) or "(none)"
        raise ValueError(
            f"No workbook config for {marketplace_key}. Known keys: {known}. Edit {_CONFIG_PATH}."
        )
    return MarketplaceWorkbookConfig.model_validate(entry)


def workbook_layout_for_key(
    marketplace_key: str,
    *,
    sheet_name: str | None = None,
    header_label_row: int | None = None,
    machine_key_row: int | None = None,
    data_start_row: int | None = None,
) -> WorkbookLayout:
    cfg = config_for_key(marketplace_key)
    return WorkbookLayout(
        sheet_name=cfg.sheet_name if sheet_name is None else sheet_name,
        header_label_row=(cfg.header_label_row if header_label_row is None else header_label_row),
        machine_key_row=cfg.machine_key_row if machine_key_row is None else machine_key_row,
        data_start_row=cfg.data_start_row if data_start_row is None else data_start_row,
        enum_discovery=cfg.enum_discovery,
        valid_values_sheet=cfg.valid_values_sheet,
        dropdown_lists_sheet=cfg.dropdown_lists_sheet,
        data_definitions_sheet=cfg.data_definitions_sheet,
    )


def listing_upload_filename(original: str | None, kind: WorkbookKind) -> str:
    """Keep the uploaded stem; extension always matches detected Open XML kind."""
    name = Path(original or "").name.strip()
    stem = Path(name).stem if name else "listing-template"
    if not stem or stem.startswith("."):
        stem = "listing-template"
    return f"{stem}.{kind}"


def metadata_for_listing_upload(
    *,
    marketplace_name: str,
    content: bytes,
    original_filename: str | None,
    existing_metadata: dict | None = None,
) -> ListingTemplateMetadata:
    """Build listing_template.metadata from the marketplace adapter + uploaded file.

    Row offsets always come from the adapter (Flipkart is 1/2/5, not Amazon 4/5/7).
    ``sheet_name`` is kept when it already exists on the row *and* in this workbook
    (category sheets like ``curtain``). Otherwise the adapter default is used.
    """
    kind = detect_openxml_kind(content)
    layout = workbook_layout_for_key(marketplace_key_from_name(marketplace_name))
    sheet_names = _sheet_names(content, kind=kind)
    sheet_name = layout.sheet_name
    existing_sheet = ""
    if isinstance(existing_metadata, dict):
        existing_sheet = str(existing_metadata.get("sheet_name") or "").strip()
    if existing_sheet and existing_sheet in sheet_names:
        sheet_name = existing_sheet
    return ListingTemplateMetadata(
        filename=listing_upload_filename(original_filename, kind),
        sheet_name=sheet_name,
        header_label_row=layout.header_label_row,
        machine_key_row=layout.machine_key_row,
        data_start_row=layout.data_start_row,
    )


def clear_config_cache() -> None:
    _load_raw.cache_clear()


def _sheet_names(content: bytes, *, kind: WorkbookKind) -> list[str]:
    workbook = load_workbook(
        io.BytesIO(content),
        read_only=True,
        data_only=True,
        keep_vba=kind == "xlsm",
    )
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()
