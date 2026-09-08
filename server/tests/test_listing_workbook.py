"""Fill listing workbooks for .xlsx / .xlsm templates."""

from __future__ import annotations

import io
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from dto.listing_config import ListingTemplateMetadata
from utils.listing_workbook import fill_workbook


def _xlsx_bytes(*, sheet_name: str, header: str) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.cell(4, 1, header)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_fill_xlsx_writes_openxml() -> None:
    metadata = ListingTemplateMetadata(
        filename="listing-template.xlsx",
        sheet_name="Template",
        data_start_row=7,
    )
    template = _xlsx_bytes(sheet_name="Template", header="item_sku")

    filled = fill_workbook(template, metadata=metadata, rows=[{1: "SKU-1"}])

    assert filled.filename == "listing-template_filled.xlsx"
    assert filled.content_type == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert filled.content.startswith(b"PK")

    workbook = load_workbook(io.BytesIO(filled.content))
    assert workbook["Template"].cell(7, 1).value == "SKU-1"


def test_fill_xlsx_does_not_mark_file_as_macro_enabled() -> None:
    metadata = ListingTemplateMetadata(
        filename="Myntra Listing File - Bedsheet.xlsx",
        sheet_name="Bedsheets",
        header_label_row=3,
        machine_key_row=3,
        data_start_row=4,
    )
    template = _xlsx_bytes(sheet_name="Bedsheets", header="styleid")
    filled = fill_workbook(template, metadata=metadata, rows=[{1: "SKU-1"}])

    assert filled.filename.endswith(".xlsx")
    with ZipFile(io.BytesIO(filled.content)) as archive:
        content_types = archive.read("[Content_Types].xml").decode()
        rels = archive.read("xl/_rels/workbook.xml.rels").decode()
    assert "macroEnabled" not in content_types
    assert "vbaProject" not in rels


def test_fill_rejects_xls() -> None:
    metadata = ListingTemplateMetadata(filename="flipkart.xls", sheet_name="bedsheet")
    with pytest.raises(ValueError, match="must be .xlsx or .xlsm"):
        fill_workbook(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest", metadata=metadata, rows=[])


def test_fill_rejects_non_excel_bytes() -> None:
    metadata = ListingTemplateMetadata(filename="listing.xlsx", sheet_name="bedsheet")
    with pytest.raises(ValueError, match="not an Excel workbook"):
        fill_workbook(b"not-excel", metadata=metadata, rows=[])


def test_fill_missing_sheet() -> None:
    metadata = ListingTemplateMetadata(
        filename="flipkart.xlsx",
        sheet_name="bedsheet",
        data_start_row=5,
    )
    template = _xlsx_bytes(sheet_name="other", header="SKU")
    with pytest.raises(ValueError, match="Sheet 'bedsheet' not found"):
        fill_workbook(template, metadata=metadata, rows=[{1: "x"}])
